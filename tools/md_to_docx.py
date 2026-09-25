"""Μετατρέπει τα docs/*.md σε .docx (για το Copilot, που διαβάζει Word).

    python tools/md_to_docx.py docs/RULES_V3.1_Panduit_Τιμές.md
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def inline(par, text):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = par.add_run(part[1:-1])
            run.font.name = 'Consolas'
        else:
            par.add_run(part)


def convert(src, dst):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    lines = open(src, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s\-|]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Light Grid Accent 1'
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    par = table.cell(r, c).paragraphs[0]
                    inline(par, text)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        if line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                run = doc.add_paragraph().add_run(lines[i])
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                i += 1
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('> '):
            par = doc.add_paragraph()
            inline(par, line[2:])
            for run in par.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x80, 0x40, 0x00)
        elif re.match(r'^\s*- ', line):
            depth = (len(line) - len(line.lstrip())) // 2
            par = doc.add_paragraph(style='List Bullet' if depth == 0 else 'List Bullet 2')
            inline(par, line.strip()[2:])
        elif re.match(r'^\d+\. ', line):
            par = doc.add_paragraph(style='List Number')
            inline(par, re.sub(r'^\d+\. ', '', line))
        elif line.strip() == '---':
            pass
        elif line.strip():
            par = doc.add_paragraph()
            inline(par, line)
            if line.startswith('**') and line.endswith('**') and '=' in line:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
    doc.save(dst)


if __name__ == '__main__':
    for src in sys.argv[1:]:
        convert(src, src[:-3] + '.docx')
        print(src[:-3] + '.docx')
